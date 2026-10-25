#!/usr/bin/env python3
"""Export researched RYM profiles without a completed catalog link to Word.

Reads preserved research and a read-only native snapshot. No catalog writes.
Requires python-docx; output must be in SSD staging before publication.
"""
import argparse
import collections
import datetime
import hashlib
import json
import re
from pathlib import Path
import shutil
import unicodedata
import urllib.parse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
STORE = Path('/mnt/1tb/Shared/coleccionistas-rock-venezolano')
SOURCE = STORE / 'reports'
GROUPS = [
    ('documented', 'Identidad documentada; vínculo sin resolver'),
    ('earlier', 'Propuestas anteriores con citas; validación final pendiente'),
    ('partial', 'Evidencia parcial; identificación sin cerrar'),
    ('distinct', 'Homónimos: el candidato comparado es otra identidad'),
    ('editorial', 'Anexo: referencias conservadas fuera del catálogo principal'),
]
LABELS = {
    'documented': 'Identidad externa documentada; no hay vínculo seguro con CRV.',
    'earlier': 'Propuesta anterior con citas comprobadas; no tiene validación nativa final.',
    'partial': 'Hay documentación, pero faltan pruebas para confirmar el vínculo.',
    'distinct': 'Se descartó este candidato CRV por ser otra identidad; no se descarta el perfil RYM.',
    'editorial': 'Identidad de referencia; exclusión editorial del catálogo principal.',
}
XML_BAD = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(text):
    return XML_BAD.sub('', ' '.join(str(text or '').split()))


def sortkey(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', value).casefold() if not unicodedata.combining(c))


def jsonl(file):
    return [json.loads(line) for line in file.read_text().splitlines() if line.strip()]


def clipped(text, limit=390):
    text = clean(text)
    if len(text) <= limit:
        return text
    head = text[:limit]
    sentence = re.search(r'^(.+[.!?])\s', head)
    if sentence and len(sentence[1]) >= limit // 2:
        return sentence[1]
    return head.rsplit(' ', 1)[0].rstrip(' ,;:') + '…'


def link(paragraph, label, target):
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '176589'); properties.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); properties.append(underline)
    run.append(properties)
    text = OxmlElement('w:t'); text.text = clean(label); run.append(text)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def page_number(paragraph):
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'
    run.append(text); field.append(run); paragraph._p.append(field)


def quote_excerpt(quote, budget):
    words = clean(quote).split()
    taken = []
    for word in words[:min(22, budget)]:
        if len(' '.join([*taken, word])) > 210:
            break
        taken.append(word)
    return ' '.join(taken), len(taken)


def keep_profiles_together(document):
    block = []

    def finish():
        for i, paragraph in enumerate(block):
            paragraph.paragraph_format.keep_with_next = i < len(block) - 1
            paragraph.paragraph_format.widow_control = True
        block.clear()

    for paragraph in document.paragraphs:
        boundary = paragraph.style.name in ('Heading 1', 'Heading 2') or paragraph._p.xpath('.//w:br[@w:type="page"]')
        if boundary:
            finish()
        if paragraph.style.name == 'Heading 2' or block:
            block.append(paragraph)
    finish()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    out = Path(args.out)
    live = json.loads((out / 'catalogo-consultado.json').read_text())
    profiles = {r['href']: r for r in jsonl(SOURCE / 'rym-integral-2026-10-21/corpus/artistas.jsonl')}
    ledger = json.loads((SOURCE / 'rym-integral-2026-10-21/identity-ledger.json').read_text())['identities']
    active = {r['id']: r['name'] for r in live['activeArtists']}
    existing = {r['href'] for r in ledger if r['status'] == 'resolved' and r['artistId'] in active}
    native = {r['value']['url'].replace('https://rateyourmusic.com', '').rstrip('/') for r in live['nativeRymLinks']}
    prior_file = SOURCE / 'rym-hermes-todos-2026-10-22/resultados-verificados.jsonl'
    deep_file = SOURCE / 'rym-verificacion-profunda-2026-10-22/resultados-finales.jsonl'
    rows = {r['rymHref']: {**r, '_origin': 'primera investigación'} for r in jsonl(prior_file)}
    rows.update({r['rymHref']: {**r, '_origin': 'verificación profunda'} for r in jsonl(deep_file)})
    people = {r['id']: r['name'] for r in json.loads((SOURCE / 'rym-verificacion-profunda-2026-10-22/catalogo-lectura.json').read_text())['persons']}
    raw_hashes = {}
    extraction_cache = {}
    excluded = collections.Counter()
    issues = []
    entries = []

    def intact(check, evidence):
        path = Path(check.get('sourcePath') or '')
        prefix = '/home/brian/crv-staging/'
        if str(path).startswith(prefix):
            path = STORE / str(path)[len(prefix):]
        expected = check.get('sourceSha256')
        if not expected or not path.is_file():
            return None
        key = str(path)
        if key not in raw_hashes:
            raw_hashes[key] = hashlib.sha256(path.read_bytes()).hexdigest()
        if raw_hashes[key] != expected:
            # A Tavily response contains several extracted documents. Its
            # historical SHA belongs to raw_content, not the JSON envelope.
            if check.get('verificationSource') != 'independent_tavily_extraction':
                return None
            if key not in extraction_cache:
                extraction_cache[key] = json.loads(path.read_text())
            canon = lambda u: urllib.parse.unquote(urllib.parse.urlsplit(u).netloc.removeprefix('www.') + urllib.parse.urlsplit(u).path.rstrip('/') + '?' + urllib.parse.urlsplit(u).query)
            for item in extraction_cache[key].get('results', []):
                content = item.get('raw_content') or ''
                if canon(item.get('url', '')) != canon(evidence['url']):
                    continue
                if hashlib.sha256(content.encode()).hexdigest() != expected:
                    continue
                if clean(unicodedata.normalize('NFKC', evidence['quote'])) not in clean(unicodedata.normalize('NFKC', content)):
                    continue
                return {'sourcePath': str(path), 'sourceSha256': expected,
                    'sourceContainerSha256': raw_hashes[key], 'hashScope': 'extracted_document',
                    'extractedDocumentUrl': item['url']}
            return None
        return {'sourcePath': str(path), 'sourceSha256': expected, 'hashScope': 'file'}

    for index, (href, row) in enumerate(rows.items(), 1):
        if href in existing or href in native:
            excluded['already_linked'] += 1
            continue
        status = row.get('finalStatus')
        group = {'documented_identity_without_catalog_link': 'documented',
                 'unresolved_with_specific_gap': 'partial',
                 'verified_distinct_candidate': 'distinct',
                 'verified_existing_editorial_reference': 'editorial'}.get(status, 'earlier')
        checks = row.get('citationChecks', []) if status else row.get('independentCitationChecks', [])
        good = []
        for evidence, check in zip(row.get('evidence', []), checks, strict=True):
            url = evidence.get('url') or ''
            parsed = urllib.parse.urlsplit(url)
            if parsed.scheme not in ('http', 'https') or not parsed.netloc:
                continue
            if not check.get('quoteFound') or not (check.get('nonRym') or check.get('nonRymSource')):
                continue
            if any(s in parsed.netloc.lower() for s in ('myspace', 'last.fm', 'lastfm', 'rateyourmusic')) or '/reviews/' in parsed.path.lower():
                continue
            # Candidate-only citations cannot demonstrate anything about a RYM
            # identity that has not been resolved. References already have a
            # final semantic acceptance and keep their documentary citations.
            if evidence.get('side') == 'candidate' and group != 'editorial':
                continue
            good.append((evidence, check))
        good.sort(key=lambda pair: (not pair[1].get('primaryEligible', False), pair[0].get('side') == 'candidate', len(pair[0].get('quote') or '')))
        selected = []
        seen_urls = set()
        for evidence, check in good:
            url = evidence['url']
            if url in seen_urls:
                continue
            pin = intact(check, evidence)
            if not pin:
                issues.append({'rymHref': href, 'url': url, 'reason': 'original_missing_or_sha_changed'})
                continue
            seen_urls.add(url)
            selected.append({**evidence, 'quote': clean(evidence['quote']),
                **pin,
                'primaryEligible': bool(check.get('primaryEligible')), 'citationVerified': True})
            if len(selected) == 2:
                break
        if not selected:
            excluded['without_verified_external_profile_evidence'] += 1
            continue
        profile = profiles.get(href, {})
        reason = row.get('finalReason') or row.get('semanticReview', {}).get('reason') or row.get('reason') or ''
        blocking = row.get('semanticReview', {}).get('blockingIssues') or []
        aid, pid = row.get('candidateArtistId'), row.get('candidatePersonId')
        entries.append({'rymHref': href, 'rymUrl': f'https://rateyourmusic.com{href}',
            'name': profile.get('nombre') or href.rsplit('/', 1)[-1], 'group': group,
            'candidateArtistId': aid, 'candidateArtistName': active.get(aid),
            'candidatePersonId': pid, 'candidatePersonName': people.get(pid),
            'status': LABELS[group], 'reason': reason, 'blockingIssues': blocking,
            'nextCheck': row.get('nextCheck'), 'contradictions': row.get('contradictions') or [],
            'evidence': selected, 'origin': row['_origin']})
        if index % 250 == 0:
            print(f'Expedientes cotejados: {index}/{len(rows)}', flush=True)

    entries.sort(key=lambda e: (next(i for i, (k, _) in enumerate(GROUPS) if k == e['group']), sortkey(e['name']), e['rymHref']))
    assert len({e['rymHref'] for e in entries}) == len(entries)
    assert not any(e['rymHref'] in native or e['rymHref'] in existing for e in entries)
    counts = collections.Counter(e['group'] for e in entries)
    main_count = len(entries) - counts['editorial']
    now = datetime.datetime.now(datetime.timezone.utc)
    summary = {'createdAt': now.isoformat(), 'corpusProfiles': len(profiles),
        'previouslyResolvedProfiles': len(existing), 'newlyAppliedProfiles': len(native),
        'profilesWithoutCompletedLink': len(profiles) - len(existing | native),
        'documentMainEntries': main_count, 'editorialAppendixEntries': counts['editorial'],
        'documentEntries': len(entries), 'groups': dict(counts), 'excluded': dict(excluded),
        'delegatedProfilesOutsideLocalFinalVerification': len(profiles) - len(existing | native) - (len(rows) - excluded['already_linked']),
        'distinctSourceFilesRehashed': len(raw_hashes), 'sourceIntegrityIssues': issues,
        'scope': 'Perfiles RYM sin vínculo final, con al menos una cita externa conservada y comprobada. No afirma que el artista no exista por nombre en CRV. Las referencias editoriales figuran en anexo.',
        'sourcePins': {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in (prior_file, deep_file)}}
    (out / 'artistas-con-evidencia.json').write_text(json.dumps(entries, ensure_ascii=False, indent=2) + '\n')
    (out / 'resumen.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.65), Cm(1.6)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.02
    for style_name, size, color in [('Title', 29, '17374D'), ('Heading 1', 17, '17374D'), ('Heading 2', 11, '17374D')]:
        style = doc.styles[style_name]; style.font.name = 'Calibri'; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(2)
    header = section.header.paragraphs[0]
    header.add_run('CRV  /  INVESTIGACIÓN RYM').font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run('Perfiles sin vínculo final · ' + now.strftime('%d/%m/%Y') + '  |  Página ').font.size = Pt(8)
    page_number(footer)
    doc.core_properties.title = 'Artistas RYM sin incorporación cerrada — evidencia y enlaces'
    doc.core_properties.subject = 'Perfiles sin vínculo final en CRV, con evidencia externa conservada'
    doc.core_properties.author = 'Coleccionistas de Rock Venezolano'
    doc.core_properties.language = 'es-VE'

    doc.add_paragraph('Artistas RYM\nsin incorporación cerrada', 'Title')
    doc.add_paragraph('Enlaces, evidencia disponible y motivos de retención', 'Subtitle')
    doc.add_paragraph(f'{main_count:,} perfiles en la lista principal · {counts["editorial"]:,} referencias en anexo'.replace(',', '.'))
    doc.add_paragraph('Corte de consulta: ' + now.strftime('%d/%m/%Y, %H:%M UTC'))
    doc.add_heading('Qué contiene esta lista', 1)
    doc.add_paragraph('Se incluyen los perfiles de RYM cuya incorporación o vinculación segura a CRV no quedó cerrada y que tienen al menos una cita externa localizada, comprobada y conservada. Se consideran la primera investigación y la revisión profunda posterior; prevalece siempre la decisión más reciente.')
    doc.add_paragraph('Una ficha candidata puede existir ya en CRV. Su aparición aquí significa que el vínculo concreto con el perfil RYM sigue sin validación final; no demuestra que el músico esté ausente del catálogo. Los identificadores de candidatos son contexto, nunca enlaces aprobados.')
    doc.add_paragraph('Quedan fuera de la lista los perfiles ya resueltos y los 70 vinculados en la última aplicación. Tampoco se rellenan evidencias que faltan: los perfiles sin una cita externa comprobada sobre su identidad no entran en la lista principal. MySpace, Last.fm y las reviews de RYM no se utilizan.')
    doc.add_paragraph('Una cita comprobada acredita que el fragmento aparece en la fuente conservada. No basta por sí sola para confirmar una identidad, una relación con CRV o todos los datos del artista. La condición pendiente y las contradicciones siguen vigentes.')
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    table.rows[0].cells[0].text = 'Grupo'; table.rows[0].cells[1].text = 'Perfiles'
    for key, title in GROUPS:
        cells = table.add_row().cells; cells[0].text = title; cells[1].text = str(counts[key])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_heading('Cómo leer las fichas', 1)
    doc.add_paragraph('Cada ficha incluye nombre, enlace RYM clicable, situación, motivo y hasta dos fuentes externas. Las citas son fragmentos breves; las fuentes repetidas pueden mostrarse con un resumen del hallazgo. El expediente JSON complementario conserva los motivos completos, las citas y las huellas SHA-256 de los originales.')
    doc.add_paragraph('La lista principal se ordena por situación y alfabéticamente dentro de cada grupo. El anexo editorial recoge identidades que se conservan fuera del catálogo principal por una decisión vigente: no son altas pendientes.')
    doc.add_paragraph(f'Cobertura de control: {len(profiles):,} perfiles en el corpus; {len(existing):,} resueltos previamente; {len(native):,} vinculados en la última aplicación. Consulta de catálogo en solo lectura; no se crearon claims ni se modificaron fichas.'.replace(',', '.'))
    doc.add_page_break()

    budgets = collections.Counter()
    number = 0
    for group, title in GROUPS:
        if not counts[group]:
            continue
        if group != GROUPS[0][0]:
            doc.add_page_break()
        doc.add_heading(title, 1)
        doc.add_paragraph(f'{counts[group]} perfiles. ' + LABELS[group])
        if group == 'earlier':
            doc.add_paragraph('Estas propuestas anteriores tenían citas comprobadas y no entraron en el bloque de 1.520 verificaciones profundas. No se presentan como aprobaciones: falta su validación nativa y, cuando corresponda, completar los dos ejes de identidad, independencia y alcance editorial.')
        if group == 'distinct':
            doc.add_paragraph('Se documentó que el perfil RYM y el candidato comparado son identidades distintas. Esto impide esa asociación, pero no confirma por sí solo una nueva alta ni la pertenencia del artista al alcance venezolano del catálogo.')
        for entry in (e for e in entries if e['group'] == group):
            number += 1
            doc.add_paragraph(f'{number:04d} · {entry["name"]}', 'Heading 2')
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run('RYM: ').bold = True
            link(p, entry['rymUrl'], entry['rymUrl'])
            candidates = []
            if entry['candidateArtistId']:
                candidates.append(f'artista {entry["candidateArtistId"]}' + (f' ({entry["candidateArtistName"]})' if entry['candidateArtistName'] else ''))
            if entry['candidatePersonId']:
                candidates.append(f'persona {entry["candidatePersonId"]}' + (f' ({entry["candidatePersonName"]})' if entry['candidatePersonName'] else ''))
            if candidates:
                p = doc.add_paragraph('Candidato CRV sin enlace aprobado: ' + '; '.join(candidates))
                if group == 'editorial':
                    p.text = 'Referencia identificada: ' + '; '.join(candidates)
                for run in p.runs:
                    run.font.size = Pt(8.5)
            p = doc.add_paragraph()
            p.add_run('Motivo: ').bold = True
            p.add_run(clipped(entry['reason'], 430))
            if entry['blockingIssues']:
                p = doc.add_paragraph()
                p.add_run('Carencia concreta: ').bold = True
                p.add_run(clipped(' · '.join(entry['blockingIssues']), 280))
            elif entry['nextCheck'] and group != 'editorial':
                p = doc.add_paragraph()
                p.add_run('Carencia concreta: ').bold = True
                p.add_run(clipped(entry['nextCheck'], 280))
            for i, evidence in enumerate(entry['evidence'], 1):
                p = doc.add_paragraph()
                p.add_run(f'Evidencia {i}: ').bold = True
                host = urllib.parse.urlsplit(evidence['url']).netloc.removeprefix('www.')
                link(p, host + ' — abrir fuente', evidence['url'])
                p.add_run(' · cita comprobada')
                fragment, used = quote_excerpt(evidence['quote'], max(0, 25 - budgets[evidence['url']]))
                budgets[evidence['url']] += used
                if fragment and used >= 4:
                    p.add_run('\n«' + fragment + '»').italic = True
                else:
                    p.add_run('\nHallazgo señalado: ' + clipped(evidence.get('supports') or 'La fuente conserva el fragmento documentado en el expediente.', 180))
            if entry['contradictions']:
                p = doc.add_paragraph()
                p.add_run('Discrepancia: ').bold = True
                p.add_run(clipped(entry['contradictions'][0], 220))
    assert number == len(entries)
    assert all(n <= 25 for n in budgets.values())
    keep_profiles_together(doc)
    filename = 'Artistas-RYM-no-incorporados-con-evidencias.docx'
    doc.save(out / filename)
    # Native package checks before rendering: every profile has a distinct RYM
    # hyperlink and an entry, and the document can be reopened as Word.
    reopened = Document(out / filename)
    targets = [r.target_ref for r in reopened.part.rels.values() if r.reltype == RT.HYPERLINK]
    assert {e['rymUrl'] for e in entries}.issubset(set(targets))
    assert sum(1 for p in reopened.paragraphs if p.style.name == 'Heading 2') == len(entries)
    summary['docxFile'] = filename
    summary['docxSha256'] = hashlib.sha256((out / filename).read_bytes()).hexdigest()
    summary['verifiedRymHyperlinks'] = len(entries)
    (out / 'resumen.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    shutil.copyfile(Path(__file__), out / 'metodo-exportacion.py')
    print(json.dumps({k: v for k, v in summary.items() if k not in ('sourcePins', 'sourceIntegrityIssues')}
        | {'sourceIntegrityIssueCount': len(issues)}, ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__':
    main()
